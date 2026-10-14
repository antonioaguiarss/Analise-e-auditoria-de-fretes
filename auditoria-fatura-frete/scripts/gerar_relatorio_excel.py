"""
Gera resultados/relatorio_auditoria.xlsx a partir dos CSVs produzidos pelo SQL.

    python scripts/montar_banco.py          # antes: gera os CSVs de resultado
    python scripts/gerar_relatorio_excel.py

A planilha repete a regra do SQL com FÓRMULAS (valor de tabela, diferença, status),
então dá para mudar a tolerância (célula azul) e ver o resumo recalcular.
"""
import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
RES = RAIZ / "resultados"
FONTE = "Arial"

f_normal = Font(name=FONTE, size=10)
f_negrito = Font(name=FONTE, size=10, bold=True)
f_titulo = Font(name=FONTE, size=14, bold=True)
f_nota = Font(name=FONTE, size=9, italic=True, color="595959")
f_input = Font(name=FONTE, size=10, color="0000FF", bold=True)
f_cab = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
preench_cab = PatternFill("solid", start_color="1F3864")
preench_input = PatternFill("solid", start_color="FFFF00")
preench_kpi = PatternFill("solid", start_color="F2F2F2")
borda = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

RS = '"R$" #,##0.00;[Red]-"R$" #,##0.00;"-"'
PCT = "0.0%"


def ler(nome):
    with open(RES / f"{nome}.csv", newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh, delimiter=";"))


def cabecalho(ws, linha, titulos, col_ini=1):
    for i, t in enumerate(titulos):
        c = ws.cell(row=linha, column=col_ini + i, value=t)
        c.font, c.fill, c.border = f_cab, preench_cab, borda
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def num(txt):
    return float(txt) if txt not in ("", None) else None


detalhe = ler("auditoria_detalhada")
duplic = ler("duplicidades")
N = len(detalhe)
ULT = N + 1  # última linha de dados na aba Detalhe

wb = Workbook()
ws_res = wb.active
ws_res.title = "Resumo"
ws_det = wb.create_sheet("Detalhe")
ws_dup = wb.create_sheet("Duplicidades")

# ---------------------------------------------------------------- Detalhe
cols = ["Nº fatura", "CT-e", "Transportadora", "Data emissão", "Mês (nº)", "Rota", "Origem", "Destino",
        "Peso (kg)", "Tarifa (R$/kg)", "Frete mínimo (R$)", "Valor cobrado (R$)", "Valor tabela (R$)",
        "Diferença (R$)", "Diferença (%)", "Status"]
cabecalho(ws_det, 1, cols)
ws_det.row_dimensions[1].height = 32
for i, d in enumerate(detalhe, start=2):
    ws_det.cell(row=i, column=1, value=d["numero_fatura"])
    ws_det.cell(row=i, column=2, value=d["cte"])
    ws_det.cell(row=i, column=3, value=d["transportadora"])
    ws_det.cell(row=i, column=4, value=datetime.strptime(d["data_emissao"], "%Y-%m-%d").date())
    ws_det.cell(row=i, column=5, value=f"=MONTH(D{i})")
    ws_det.cell(row=i, column=6, value=d["rota_id"])
    ws_det.cell(row=i, column=7, value=d["origem"])
    ws_det.cell(row=i, column=8, value=d["destino"])
    ws_det.cell(row=i, column=9, value=num(d["peso_kg"]))
    ws_det.cell(row=i, column=10, value=num(d["tarifa_kg"]))
    ws_det.cell(row=i, column=11, value=num(d["frete_minimo"]))
    ws_det.cell(row=i, column=12, value=num(d["valor_cobrado"]))
    # Regra do contrato: o MAIOR entre (peso x tarifa) e o frete mínimo
    ws_det.cell(row=i, column=13, value=f'=IF(J{i}="","",ROUND(MAX(I{i}*J{i},K{i}),2))')
    ws_det.cell(row=i, column=14, value=f'=IF(M{i}="","",ROUND(L{i}-M{i},2))')
    ws_det.cell(row=i, column=15, value=f'=IF(M{i}="","",N{i}/M{i})')
    ws_det.cell(row=i, column=16,
                value=(f'=IF(M{i}="","SEM TARIFA NO CONTRATO",IF(N{i}>Resumo!$B$4,"ACIMA DA TABELA",'
                       f'IF(N{i}<-Resumo!$B$4,"ABAIXO DA TABELA","OK")))'))
    for c in range(1, 17):
        ws_det.cell(row=i, column=c).font = f_normal
    ws_det.cell(row=i, column=4).number_format = "dd/mm/yyyy"
    ws_det.cell(row=i, column=9).number_format = "#,##0.0"
    ws_det.cell(row=i, column=10).number_format = "0.0000"
    for c in (11, 12, 13, 14):
        ws_det.cell(row=i, column=c).number_format = RS
    ws_det.cell(row=i, column=15).number_format = PCT

for c, larg in enumerate([13, 14, 24, 13, 9, 7, 20, 20, 11, 12, 14, 16, 16, 15, 12, 26], start=1):
    ws_det.column_dimensions[get_column_letter(c)].width = larg
ws_det.freeze_panes = "A2"
ws_det.auto_filter.ref = f"A1:P{ULT}"
ws_det.conditional_formatting.add(
    f"A2:P{ULT}", FormulaRule(formula=['$P2="ACIMA DA TABELA"'], fill=PatternFill("solid", bgColor="F8CBAD")))
ws_det.conditional_formatting.add(
    f"A2:P{ULT}", FormulaRule(formula=['$P2="SEM TARIFA NO CONTRATO"'], fill=PatternFill("solid", bgColor="FFE699")))


def rng(letra):
    return f"Detalhe!${letra}$2:${letra}${ULT}"


STATUS, DIF, COBR, TAB, TRANSP, MES = rng("P"), rng("N"), rng("L"), rng("M"), rng("C"), rng("E")

# ---------------------------------------------------------------- Duplicidades
cabecalho(ws_dup, 1, ["CT-e", "Qtd. faturas", "Faturas", "Transportadora", "Total cobrado (R$)", "Valor a mais (R$)"])
for i, d in enumerate(duplic, start=2):
    vals = [d["cte"], int(d["qtd_faturas"]), d["faturas"], d["transportadora"],
            num(d["total_cobrado"]), num(d["valor_a_mais"])]
    for c, v in enumerate(vals, start=1):
        cel = ws_dup.cell(row=i, column=c, value=v)
        cel.font = f_normal
        if c in (5, 6):
            cel.number_format = RS
DUP_ULT = len(duplic) + 1
for c, larg in enumerate([14, 12, 24, 24, 18, 18], start=1):
    ws_dup.column_dimensions[get_column_letter(c)].width = larg
ws_dup.cell(row=DUP_ULT + 2, column=1,
            value="Fonte: view vw_duplicidades (sql/02_auditoria.sql). 'Valor a mais' = total cobrado menos uma cobrança.").font = f_nota

# ---------------------------------------------------------------- Resumo
ws_res["A1"] = "Auditoria de fatura de frete - Resumo"
ws_res["A1"].font = f_titulo
ws_res["A2"] = "Base 100% fictícia (jan-set/2026). Mesma regra do SQL, escrita em fórmulas."
ws_res["A2"].font = f_nota

ws_res["A4"] = "Tolerância de arredondamento (R$)"
ws_res["B4"] = 1.0
ws_res["B4"].font, ws_res["B4"].fill, ws_res["B4"].number_format = f_input, preench_input, RS
ws_res["C4"] = "<- célula editável: diferenças até esse valor contam como OK"
ws_res["C4"].font = f_nota
ws_res["A4"].font = f_negrito

kpis = [
    ("Faturas analisadas", f"=COUNTA({rng('A')})", "#,##0"),
    ("Faturas acima da tabela", f'=COUNTIF({STATUS},"ACIMA DA TABELA")', "#,##0"),
    ("Valor cobrado a mais (R$)", f'=SUMIFS({DIF},{STATUS},"ACIMA DA TABELA")', RS),
    ("Excesso sobre o valor de tabela (%)", f'=B9/SUM({TAB})', "0.00%"),
    ("Fretes em rota sem tarifa no contrato", f'=COUNTIF({STATUS},"SEM TARIFA NO CONTRATO")', "#,##0"),
    ("Valor desses fretes sem tarifa (R$)", f'=SUMIFS({COBR},{STATUS},"SEM TARIFA NO CONTRATO")', RS),
    ("Cobranças em duplicidade (R$ a mais)", f"=SUM(Duplicidades!$F$2:$F${DUP_ULT})", RS),
    ("Potencial a recuperar (R$)", "=B9+B13", RS),
]
ws_res["A6"], ws_res["B6"] = "Indicador", "Valor"
cabecalho(ws_res, 6, ["Indicador", "Valor"])
for i, (nome, formula, fmt) in enumerate(kpis, start=7):
    a, b = ws_res.cell(row=i, column=1, value=nome), ws_res.cell(row=i, column=2, value=formula)
    a.font, b.font = f_normal, f_negrito
    a.fill = b.fill = preench_kpi
    a.border = b.border = borda
    b.number_format = fmt
ws_res["A15"].font = f_nota
ws_res["A15"] = "Potencial a recuperar = excesso acima da tabela + duplicidades. Nesta base as duplicadas estão dentro da tabela, então não há dupla contagem."

# Por transportadora
ws_res["A17"] = "Por transportadora"
ws_res["A17"].font = f_negrito
cabecalho(ws_res, 18, ["Transportadora", "Faturas", "Acima da tabela", "% acima", "Excesso (R$)"])
ordem = [r["transportadora"] for r in ler("por_transportadora")]  # já vem ordenado por excesso
for i, t in enumerate(ordem, start=19):
    ws_res.cell(row=i, column=1, value=t)
    ws_res.cell(row=i, column=2, value=f'=COUNTIFS({TRANSP},A{i},{STATUS},"<>SEM TARIFA NO CONTRATO")')
    ws_res.cell(row=i, column=3, value=f'=COUNTIFS({TRANSP},A{i},{STATUS},"ACIMA DA TABELA")')
    ws_res.cell(row=i, column=4, value=f"=IF(B{i}=0,0,C{i}/B{i})")
    ws_res.cell(row=i, column=5, value=f'=SUMIFS({DIF},{TRANSP},A{i},{STATUS},"ACIMA DA TABELA")')
    for c in range(1, 6):
        ws_res.cell(row=i, column=c).font = f_normal
        ws_res.cell(row=i, column=c).border = borda
    ws_res.cell(row=i, column=4).number_format = PCT
    ws_res.cell(row=i, column=5).number_format = RS
fim_t = 18 + len(ordem)

# Por mês
ini_m = fim_t + 3
ws_res.cell(row=ini_m - 1, column=1, value="Por mês de emissão (2026)").font = f_negrito
cabecalho(ws_res, ini_m, ["Mês (nº)", "Faturas", "Acima da tabela", "Excesso (R$)"])
meses = sorted({int(r["mes"][5:7]) for r in ler("por_mes")})
for i, m in enumerate(meses, start=ini_m + 1):
    ws_res.cell(row=i, column=1, value=m)
    ws_res.cell(row=i, column=2, value=f'=COUNTIFS({MES},A{i},{STATUS},"<>SEM TARIFA NO CONTRATO")')
    ws_res.cell(row=i, column=3, value=f'=COUNTIFS({MES},A{i},{STATUS},"ACIMA DA TABELA")')
    ws_res.cell(row=i, column=4, value=f'=SUMIFS({DIF},{MES},A{i},{STATUS},"ACIMA DA TABELA")')
    for c in range(1, 5):
        ws_res.cell(row=i, column=c).font = f_normal
        ws_res.cell(row=i, column=c).border = borda
    ws_res.cell(row=i, column=4).number_format = RS

for c, larg in enumerate([44, 18, 18, 14, 18], start=1):
    ws_res.column_dimensions[get_column_letter(c)].width = larg
ws_res.sheet_view.showGridLines = False

wb.save(RES / "relatorio_auditoria.xlsx")
print("Gerado:", RES / "relatorio_auditoria.xlsx")
